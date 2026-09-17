from pathlib import Path
from typing import Any, Dict, Iterable, Mapping, Sequence

import numpy as np
import pandas as pd


EXCEL_MAX_COLUMNS = 16_384


def _python_value(value: Any) -> Any:
    if isinstance(value, np.generic):
        return value.item()
    return value


def _append_dataframe(worksheet: Any, dataframe: pd.DataFrame) -> None:
    worksheet.append([str(column) for column in dataframe.columns])
    for row in dataframe.itertuples(index=False, name=None):
        worksheet.append([_python_value(value) for value in row])


def save_proposal_evaluation_workbook(
    path: Path,
    summary: pd.DataFrame,
    proposal_details: Sequence[Mapping[str, Any]],
    predictions: Mapping[str, Dict[str, Any]],
    result_fields: Sequence[str],
    predicted_subscore_names: Sequence[str],
) -> None:
    """Write evaluation summary, all proposal scores, and all trajectory points to an XLSX file."""
    try:
        from openpyxl import Workbook
    except ImportError as error:
        raise RuntimeError(
            "Writing proposal evaluation workbooks requires openpyxl. "
            "Install the repository requirements before evaluation."
        ) from error

    workbook = Workbook(write_only=True)
    summary_sheet = workbook.create_sheet("summary")
    _append_dataframe(summary_sheet, summary)

    max_proposals = max((len(detail["results"]) for detail in proposal_details), default=0)
    score_header = ["token", "selected_proposal_idx", "best_proposal_idx"]
    for proposal_idx in range(max_proposals):
        prefix = f"proposal_{proposal_idx:03d}"
        score_header.extend(f"{prefix}_real_{field}" for field in result_fields)
        score_header.append(f"{prefix}_pred_pdm_score")
        score_header.extend(
            f"{prefix}_pred_{metric_name}" for metric_name in predicted_subscore_names
        )
    if len(score_header) > EXCEL_MAX_COLUMNS:
        raise ValueError(f"proposal_scores requires {len(score_header)} columns, exceeding Excel's limit")

    score_sheet = workbook.create_sheet("proposal_scores")
    score_sheet.append(score_header)
    for detail in proposal_details:
        row = [detail["token"], detail["selected_proposal_idx"], detail["best_proposal_idx"]]
        for proposal_idx in range(max_proposals):
            if proposal_idx < len(detail["results"]):
                result = detail["results"][proposal_idx]
                row.extend(_python_value(result[field]) for field in result_fields)
                row.append(_python_value(detail["predicted_pdm_scores"][proposal_idx]))
                row.extend(
                    _python_value(detail["predicted_subscores"][metric_name][proposal_idx])
                    for metric_name in predicted_subscore_names
                )
            else:
                row.extend([None] * (len(result_fields) + 1 + len(predicted_subscore_names)))
        score_sheet.append(row)

    max_poses = max(
        (
            len(trajectory.poses)
            for prediction in predictions.values()
            for trajectory in prediction.get("proposals", [])
        ),
        default=0,
    )
    trajectory_header = ["token", "selected_proposal_idx"]
    for proposal_idx in range(max_proposals):
        for pose_idx in range(max_poses):
            prefix = f"proposal_{proposal_idx:03d}_pose_{pose_idx:03d}"
            trajectory_header.extend((f"{prefix}_x", f"{prefix}_y", f"{prefix}_heading"))
    if len(trajectory_header) > EXCEL_MAX_COLUMNS:
        raise ValueError(
            f"proposal_trajectories requires {len(trajectory_header)} columns, exceeding Excel's limit"
        )

    trajectory_sheet = workbook.create_sheet("proposal_trajectories")
    trajectory_sheet.append(trajectory_header)
    for detail in proposal_details:
        token = detail["token"]
        prediction = predictions[token]
        row = [token, detail["selected_proposal_idx"]]
        proposals = prediction.get("proposals", [])
        for proposal_idx in range(max_proposals):
            if proposal_idx < len(proposals):
                poses = np.asarray(proposals[proposal_idx].poses)
                for pose_idx in range(max_poses):
                    if pose_idx < len(poses):
                        row.extend(_python_value(value) for value in poses[pose_idx, :3])
                    else:
                        row.extend((None, None, None))
            else:
                row.extend([None] * (max_poses * 3))
        trajectory_sheet.append(row)

    workbook.save(path)
