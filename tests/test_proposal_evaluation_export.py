from types import SimpleNamespace

import pandas as pd
from openpyxl import load_workbook

from navsim.planning.script.proposal_evaluation_export import save_proposal_evaluation_workbook


def test_save_proposal_evaluation_workbook(tmp_path) -> None:
    path = tmp_path / "evaluation.xlsx"
    summary = pd.DataFrame([{"token": "scene-a", "valid": True, "score": 0.8}])
    result_fields = ["no_at_fault_collisions", "score"]
    proposal_details = [
        {
            "token": "scene-a",
            "selected_proposal_idx": 0,
            "best_proposal_idx": 0,
            "results": [
                {"no_at_fault_collisions": 1.0, "score": 0.8},
                {"no_at_fault_collisions": 0.0, "score": 0.0},
            ],
            "predicted_pdm_scores": [0.7, 0.2],
            "predicted_subscores": {"comfort": [0.9, 0.8]},
        }
    ]
    predictions = {
        "scene-a": {
            "proposals": [
                SimpleNamespace(poses=[[1.0, 2.0, 0.1], [3.0, 4.0, 0.2]]),
                SimpleNamespace(poses=[[5.0, 6.0, 0.3], [7.0, 8.0, 0.4]]),
            ]
        }
    }

    save_proposal_evaluation_workbook(
        path=path,
        summary=summary,
        proposal_details=proposal_details,
        predictions=predictions,
        result_fields=result_fields,
        predicted_subscore_names=["comfort"],
    )

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        assert workbook.sheetnames == ["summary", "proposal_scores", "proposal_trajectories"]
        score_rows = list(workbook["proposal_scores"].iter_rows(values_only=True))
        trajectory_rows = list(workbook["proposal_trajectories"].iter_rows(values_only=True))
        assert "proposal_001_real_score" in score_rows[0]
        assert "proposal_001_pred_comfort" in score_rows[0]
        assert "proposal_001_pose_001_heading" in trajectory_rows[0]
        assert trajectory_rows[1][2:8] == (1, 2, 0.1, 3, 4, 0.2)
    finally:
        workbook.close()
